# export.py
import os, time, platform
from docx import Document
from docx.shared import Pt, Cm
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from reportlab.lib.pagesizes import A4
from reportlab.lib.units import cm
from reportlab.lib import colors
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer,
                                 Table, TableStyle)
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

MODES = {
    "zh":           "纯中文",
    "en":           "纯英文",
    "two_column":   "中英对照（左右两栏）",
    "interleaved":  "中英对照（段段交替）",
}


# ---------------- Markdown ----------------
def _md_escape(s: str) -> str:
    return s.replace("|", "\\|").replace("\n", " ").strip()


def export_md(records, path, mode):
    lines = [f"# 会议记录 {time.strftime('%Y-%m-%d %H:%M')}", ""]

    if mode == "zh":
        for r in records:
            lines.append(r["zh"]); lines.append("")

    elif mode == "en":
        for r in records:
            lines.append(r["en"]); lines.append("")

    elif mode == "two_column":
        lines.append("| English | 中文 |")
        lines.append("|:--------|:-----|")
        for r in records:
            lines.append(f"| {_md_escape(r['en'])} | {_md_escape(r['zh'])} |")

    elif mode == "interleaved":
        for r in records:
            lines.append(r["en"]); lines.append("")
            lines.append(r["zh"]); lines.append("")
            lines.append("")                # 额外空行

    with open(path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))


# ---------------- Word ----------------
def _cn_font(run, name="微软雅黑", size=11.0, bold=False):
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.bold = bold
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts"); rPr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), name)


def export_docx(records, path, mode):
    doc = Document()
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Cm(2.2)
        s.left_margin = s.right_margin = Cm(2.4)
    base = doc.styles["Normal"]
    base.font.name = "Times New Roman"
    base.font.size = Pt(11)

    if mode == "zh":
        for r in records:
            p = doc.add_paragraph()
            _cn_font(p.add_run(r["zh"]), size=11)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.line_spacing = 1.45

    elif mode == "en":
        for r in records:
            p = doc.add_paragraph(r["en"])
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.line_spacing = 1.45

    elif mode == "two_column":
        table = doc.add_table(rows=1, cols=2)
        table.autofit = True
        hdr = table.rows[0].cells
        _cn_font(hdr[0].paragraphs[0].add_run("English"), size=10.5, bold=True)
        _cn_font(hdr[1].paragraphs[0].add_run("中文"),    size=10.5, bold=True)
        for r in records:
            row = table.add_row().cells
            row[0].paragraphs[0].add_run(r["en"]).font.size = Pt(10.5)
            _cn_font(row[1].paragraphs[0].add_run(r["zh"]), size=10.5)

    elif mode == "interleaved":
        for r in records:
            p = doc.add_paragraph(); p.add_run(r["en"]).font.size = Pt(11)
            p.paragraph_format.space_after = Pt(2)
            p = doc.add_paragraph()
            _cn_font(p.add_run(r["zh"]), size=11)
            p.paragraph_format.space_after = Pt(16)

    doc.save(path)


# ---------------- PDF ----------------
def _register_cn_font():
    sysname = platform.system()
    if sysname == "Windows":
        cands = [("MSYH",   r"C:\Windows\Fonts\msyh.ttc"),
                 ("SimHei", r"C:\Windows\Fonts\simhei.ttf"),
                 ("SimSun", r"C:\Windows\Fonts\simsun.ttc")]
    elif sysname == "Darwin":
        cands = [("PingFang", "/System/Library/Fonts/PingFang.ttc"),
                 ("STHeiti",  "/System/Library/Fonts/STHeiti Light.ttc")]
    else:
        cands = [("NotoSansCJK",
                  "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc")]
    for name, path in cands:
        if not os.path.exists(path): continue
        for kwargs in ({"subfontIndex": 0}, {}):
            try:
                pdfmetrics.registerFont(TTFont(name, path, **kwargs))
                return name
            except Exception:
                continue
    return "Helvetica"


def _esc(s):
    return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def export_pdf(records, path, mode):
    cn = _register_cn_font()
    en_style = ParagraphStyle("en", fontName="Times-Roman", fontSize=10.5, leading=15)
    zh_style = ParagraphStyle("zh", fontName=cn,           fontSize=10.5, leading=15)

    doc = SimpleDocTemplate(path, pagesize=A4,
                            topMargin=2*cm, bottomMargin=2*cm,
                            leftMargin=2*cm, rightMargin=2*cm)
    story = []

    if mode == "en":
        for r in records:
            story.append(Paragraph(_esc(r["en"]), en_style))
            story.append(Spacer(1, 6))

    elif mode == "zh":
        for r in records:
            story.append(Paragraph(_esc(r["zh"]), zh_style))
            story.append(Spacer(1, 6))

    elif mode == "two_column":
        data = [[Paragraph("<b>English</b>", en_style),
                 Paragraph("<b>中文</b>",   zh_style)]]
        for r in records:
            data.append([Paragraph(_esc(r["en"]), en_style),
                         Paragraph(_esc(r["zh"]), zh_style)])
        tbl = Table(data, colWidths=[8.5*cm, 8.5*cm], repeatRows=1)
        tbl.setStyle(TableStyle([
            ("VALIGN",        (0,0), (-1,-1), "TOP"),
            ("TOPPADDING",    (0,0), (-1,-1), 6),
            ("BOTTOMPADDING", (0,0), (-1,-1), 6),
            ("LINEBELOW",     (0,0), (-1,-2), 0.3, colors.lightgrey),
            ("BACKGROUND",    (0,0), (-1,0),  colors.whitesmoke),
        ]))
        story.append(tbl)

    elif mode == "interleaved":
        for r in records:
            story.append(Paragraph(_esc(r["en"]), en_style))
            story.append(Spacer(1, 2))
            story.append(Paragraph(_esc(r["zh"]), zh_style))
            story.append(Spacer(1, 14))

    doc.build(story)


EXPORTERS = {
    "docx": export_docx,
    "pdf":  export_pdf,
    "md":   export_md,
}